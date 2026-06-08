from __future__ import annotations

import csv
import io
import re
import time
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

DOCUMENT_TYPES = {
    "Sales Invoice",
    "Purchase Invoice",
    "Payment Voucher",
    "Receipt",
    "Payroll Summary",
    "Loan / HP Statement",
    "Merchant Statement",
    "Utility Bill",
}

ACCOUNT_NAMES = {
    "4100": "Sales Revenue",
    "4120": "F&B Revenue",
    "5020": "Direct Materials",
    "6100": "Salaries & Wages",
    "6200": "Rent Expense",
    "6210": "Utilities Electricity",
    "6370": "Bank Charges & Fees",
    "6380": "Insurance Expense",
    "2700": "Term Loan Payable",
}


@dataclass(frozen=True)
class SmartRule:
    label: str
    doc_type: str
    target: str
    keywords: tuple[str, ...]
    account_code: str = ""


SMART_RULES = [
    SmartRule("Bank statement layout", "Bank Statement", "WP2", ("bank statement", "account statement", "opening balance", "closing balance")),
    SmartRule("Bank transaction columns", "Bank Statement", "WP2", ("money in", "money out", "debit", "credit", "balance")),
    SmartRule("TNB utilities", "Utility Bill", "WP1", ("tnb", "tenaga", "electricity", "utility"), "6210"),
    SmartRule("Water utility", "Utility Bill", "WP1", ("air selangor", "water bill", "utility"), "6210"),
    SmartRule("Merchant payout", "Merchant Statement", "WP1", ("grab", "foodpanda", "merchant", "payout", "settlement"), "4120"),
    SmartRule("Payroll support", "Payroll Summary", "WP1", ("payroll", "salary", "epf", "socso", "eis", "pcb"), "6100"),
    SmartRule("Loan repayment support", "Loan / HP Statement", "WP1", ("loan", "hire purchase", "principal", "interest", "instalment"), "2700"),
    SmartRule("Insurance / takaful", "Payment Voucher", "WP1", ("insurance", "takaful", "premium"), "6380"),
    SmartRule("Rent payment", "Payment Voucher", "WP1", ("rental", "rent", "landlord"), "6200"),
    SmartRule("Bank charges", "Bank Statement", "WP2", ("bank charge", "bank fee", "service charge", "charges"), "6370"),
    SmartRule("Sales invoice", "Sales Invoice", "WP1", ("sales invoice", "customer invoice", "invoice to"), "4100"),
    SmartRule("Supplier invoice", "Purchase Invoice", "WP1", ("supplier invoice", "purchase invoice", "vendor invoice", "bill from"), "5020"),
    SmartRule("Receipt", "Receipt", "WP1", ("official receipt", "receipt"), "4100"),
    SmartRule("Payment voucher", "Payment Voucher", "WP1", ("payment voucher", "pv"), "5020"),
]


def extract_intake_items(filename: str, content_type: str, data: bytes) -> list[dict[str, Any]]:
    uploaded_at = current_timestamp()
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension in {"xlsx", "xlsm"}:
        rows = extract_xlsx_rows(filename, content_type, data, uploaded_at)
        if rows:
            return rows

    if extension in {"csv", "txt", "tsv"} or content_type.startswith("text/"):
        text = decode_text(data)
        rows = extract_delimited_rows(filename, content_type, text, uploaded_at)
        if rows:
            return rows
        return [single_item(filename, content_type, len(data), uploaded_at, text)]

    if extension == "pdf" or content_type == "application/pdf":
        text, warning = extract_pdf_text(data)
        item = single_item(filename, content_type, len(data), uploaded_at, text)
        if warning:
            item["warnings"].append(warning)
        return [item]

    item = single_item(filename, content_type, len(data), uploaded_at, "")
    item["warnings"].append("File type is not directly readable yet. Import as a review row.")
    return [item]


def extract_xlsx_rows(filename: str, content_type: str, data: bytes, uploaded_at: str) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    items: list[dict[str, Any]] = []

    for sheet in workbook.worksheets:
        rows = [
            [cell_to_text(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
            if any(cell_to_text(value) for value in row)
        ]
        header_index = find_header_index(rows)
        if header_index is None:
            continue

        headers = [normalize(header) for header in rows[header_index]]
        for offset, row in enumerate(rows[header_index + 1 : header_index + 51]):
            if not any(row):
                continue
            items.append(item_from_cells(
                filename=f"{filename} {sheet.title} row {offset + 1}",
                content_type=content_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                file_size=len(data),
                uploaded_at=uploaded_at,
                headers=headers,
                cells=row,
                row_text=" ".join(row),
            ))

    return items


def extract_delimited_rows(filename: str, content_type: str, text: str, uploaded_at: str) -> list[dict[str, Any]]:
    lines = [line for line in text.splitlines() if line.strip()]
    if len(lines) < 2:
        return []

    dialect = csv.excel_tab if "\t" in lines[0] else csv.excel
    rows = [[cell.strip() for cell in row] for row in csv.reader(lines, dialect=dialect)]
    header_index = find_header_index(rows)
    if header_index is None:
        return []

    headers = [normalize(header) for header in rows[header_index]]
    return [
        item_from_cells(
            filename=f"{filename} row {index + 1}",
            content_type=content_type or "text/csv",
            file_size=len(text.encode("utf-8")),
            uploaded_at=uploaded_at,
            headers=headers,
            cells=row,
            row_text=" ".join(row),
        )
        for index, row in enumerate(rows[header_index + 1 : header_index + 51])
        if any(cell.strip() for cell in row)
    ]


def item_from_cells(
    filename: str,
    content_type: str,
    file_size: int,
    uploaded_at: str,
    headers: list[str],
    cells: list[str],
    row_text: str,
) -> dict[str, Any]:
    header_text = " ".join(headers)
    looks_like_bank = bool(re.search(r"money in|deposit|credit|money out|withdrawal|debit|bank description|balance", header_text))

    detection = {
        "detectedType": "Bank Statement",
        "confidence": "High",
        "target": "WP2",
        "evidence": ["Detected bank-style columns in uploaded file."],
        "warnings": [],
        "suggestedGlAccount": "",
    } if looks_like_bank else detect_type(f"{filename} {row_text}")

    date = cell_at(headers, cells, ["date", "transaction date"]) or extract_date(row_text)
    description = cell_at(headers, cells, ["description", "bank description", "vendor", "customer", "party"]) or clean_words(filename)
    reference = cell_at(headers, cells, ["reference", "doc ref", "document ref", "invoice", "ref"]) or extract_reference(row_text, filename)
    type_from_row = cell_at(headers, cells, ["document type", "doc type", "type"])
    amount_text = cell_at(headers, cells, ["amount", "total"])
    money_in = parse_amount(cell_at(headers, cells, ["money in", "deposit", "credit"]))
    money_out = parse_amount(cell_at(headers, cells, ["money out", "withdrawal", "debit"]))
    signed_amount, signed_direction = parse_signed_amount(amount_text)
    amount = max(signed_amount, money_in, money_out)

    if detection["target"] == "WP2" and not money_in and not money_out and amount:
      if signed_direction == "OUT":
          money_out = amount
      else:
          money_in = amount

    detected_type = type_from_row if type_from_row in DOCUMENT_TYPES or type_from_row == "Bank Statement" else detection["detectedType"]
    warnings = [
        *detection["warnings"],
        "" if date else "Date was not detected.",
        "" if amount > 0 else "Amount was not detected.",
    ]

    return {
        "id": intake_id(),
        "fileName": filename,
        "fileType": content_type or "unknown",
        "fileSize": file_size,
        "uploadedAt": uploaded_at,
        "detectedType": detected_type,
        "confidence": detection["confidence"],
        "status": "Needs Review",
        "target": detection["target"],
        "date": date,
        "reference": reference,
        "party": description,
        "amount": amount,
        "moneyIn": money_in,
        "moneyOut": money_out,
        "suggestedGlAccount": cell_at(headers, cells, ["gl account", "account"]) or detection["suggestedGlAccount"],
        "notes": cell_at(headers, cells, ["note", "notes", "remarks"]) or "Extracted from uploaded workbook/spreadsheet.",
        "evidence": detection["evidence"],
        "warnings": [warning for warning in warnings if warning],
        "rawPreview": row_text[:240],
    }


def single_item(filename: str, content_type: str, file_size: int, uploaded_at: str, text: str) -> dict[str, Any]:
    detection = detect_type(f"{filename} {text[:2000]}")
    amount = extract_largest_amount(f"{filename} {text[:2000]}")
    date = extract_date(f"{filename} {text[:2000]}")
    warnings = [
        *detection["warnings"],
        "" if date else "Date was not detected.",
        "" if amount > 0 else "Amount was not detected.",
    ]

    return {
        "id": intake_id(),
        "fileName": filename,
        "fileType": content_type or "unknown",
        "fileSize": file_size,
        "uploadedAt": uploaded_at,
        "detectedType": detection["detectedType"],
        "confidence": detection["confidence"],
        "status": "Needs Review",
        "target": detection["target"],
        "date": date,
        "reference": extract_reference(f"{filename} {text[:2000]}", filename),
        "party": clean_words(filename),
        "amount": amount,
        "moneyIn": amount if detection["target"] == "WP2" else 0,
        "moneyOut": 0,
        "suggestedGlAccount": detection["suggestedGlAccount"],
        "notes": "Extracted by source document parser." if text else "File content was not readable. Review downstream fields.",
        "evidence": detection["evidence"],
        "warnings": [warning for warning in warnings if warning],
        "rawPreview": text[:240],
    }


def detect_type(source: str) -> dict[str, Any]:
    content = normalize(source)
    ranked = []
    for rule in SMART_RULES:
        hits = [keyword for keyword in rule.keywords if keyword in content]
        if hits:
            ranked.append((len(hits), rule, hits))

    ranked.sort(key=lambda item: item[0], reverse=True)
    if not ranked:
        return {
            "detectedType": "Unknown",
            "confidence": "Low",
            "target": "WP1",
            "evidence": ["No reliable keyword pattern matched."],
            "warnings": ["Imported as a review item. Choose document type and GL downstream."],
            "suggestedGlAccount": "",
        }

    score, rule, hits = ranked[0]
    confidence = "High" if score >= 2 else "Medium"
    warnings = [] if confidence == "High" else ["Only one strong clue found. Review before posting."]
    if len(ranked) > 1 and ranked[1][0] == score and ranked[1][1].doc_type != rule.doc_type:
        warnings.append(f"Also looked like {ranked[1][1].doc_type}.")

    return {
        "detectedType": rule.doc_type,
        "confidence": confidence,
        "target": rule.target,
        "evidence": [f"Matched {rule.label}: {', '.join(hits)}"],
        "warnings": warnings,
        "suggestedGlAccount": account_label(rule.account_code),
    }


def extract_pdf_text(data: bytes) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
    except Exception:
        return "", "PDF text extractor is not installed. Review downstream fields."

    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        return "", "PDF could not be read. Review downstream fields."

    if not text.strip():
        return "", "No embedded PDF text found. Scanned PDFs need OCR/AI extraction."
    return text, ""


def find_header_index(rows: list[list[str]]) -> int | None:
    for index, row in enumerate(rows[:20]):
        header_text = " ".join(normalize(cell) for cell in row)
        if re.search(r"date|transaction date", header_text) and re.search(r"amount|money in|debit|credit|total", header_text):
            return index
    return None


def cell_at(headers: list[str], cells: list[str], candidates: list[str]) -> str:
    for index, header in enumerate(headers):
        if any(candidate in header for candidate in candidates):
            return cells[index].strip() if index < len(cells) else ""
    return ""


def parse_amount(value: str | None) -> float:
    if not value:
        return 0.0
    normalized = re.sub(r"[RMrm\s(),]", "", str(value)).replace(",", "")
    try:
        return abs(float(normalized))
    except ValueError:
        return 0.0


def parse_signed_amount(value: str | None) -> tuple[float, str]:
    if not value:
        return 0.0, "IN"
    text = str(value).strip()
    amount = parse_amount(text)
    direction = "OUT" if text.startswith("-") or ("(" in text and ")" in text) else "IN"
    return amount, direction


def extract_date(source: str) -> str:
    iso = re.search(r"\b(20\d{2})[-/.](0?[1-9]|1[0-2])[-/.](0?[1-9]|[12]\d|3[01])\b", source)
    if iso:
        return f"{iso.group(3).zfill(2)} {month_name(int(iso.group(2)))}"

    day_month = re.search(r"\b(0?[1-9]|[12]\d|3[01])\s*(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\b", source, re.I)
    if day_month:
        return f"{day_month.group(1).zfill(2)} {day_month.group(2)[:3].title()}"

    slash = re.search(r"\b(0?[1-9]|[12]\d|3[01])[-/.](0?[1-9]|1[0-2])[-/.](20\d{2}|\d{2})\b", source)
    if slash:
        return f"{slash.group(1).zfill(2)} {month_name(int(slash.group(2)))}"
    return ""


def extract_reference(source: str, fallback: str) -> str:
    reference = re.search(r"\b(?:INV|PV|OR|REC|BILL|PAY|LOAN|CHQ|CN|DN)[-\s]?[A-Z0-9-]{2,}\b", source, re.I)
    if reference:
        return re.sub(r"\s+", "-", reference.group(0)).upper()
    cleaned = clean_words(fallback).upper()
    return "-".join(cleaned.split()[:3])


def extract_largest_amount(source: str) -> float:
    matches = re.findall(r"(?:RM\s*)?\(?\d{1,3}(?:,\d{3})*(?:\.\d{2})?\)?", source, re.I)
    return max([parse_amount(match) for match in matches], default=0.0)


def normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def clean_words(value: str) -> str:
    cleaned = re.sub(r"\.[^.]+$", "", value)
    cleaned = re.sub(r"[_-]+", " ", cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


def account_label(code: str) -> str:
    return f"{code} - {ACCOUNT_NAMES[code]}" if code and code in ACCOUNT_NAMES else ""


def month_name(month: int) -> str:
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    return months[month - 1] if 1 <= month <= 12 else ""


def cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def intake_id() -> str:
    return f"INT-{time.time_ns()}"


def current_timestamp() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import load_workbook

from .memory import detect_type, looks_bank_like, normalize_upload_lane
from .normalization import (
    cell_at,
    cell_prefer,
    cell_to_text,
    clean_reference,
    current_timestamp,
    decode_text,
    extract_document_amount,
    extract_reference,
    find_header_index,
    month_name,
    normalize,
    normalize_date,
    normalize_whitespace,
    parse_amount,
    parse_signed_amount,
    split_delimited_rows,
)
from .postprocess import build_item, clean_party_value
from .smart_classification import classify_revenue_context

DAY_OR_TIME_TOKENS = {
    "mon", "monday", "tue", "tues", "tuesday", "wed", "wednesday", "thu", "thurs", "thursday",
    "fri", "friday", "sat", "saturday", "sun", "sunday", "am", "pm",
}
GENERIC_SALES_LABELS = {
    "sales", "summary", "total", "grand total", "overall", "all", "others", "misc", "miscellaneous",
}
COMPANY_SUFFIXES = (
    "sdn bhd",
    "bhd",
    "enterprise",
    "trading",
    "supplies",
    "supply",
    "resources",
    "services",
    "holdings",
    "restaurant",
    "cafe",
    "kitchen",
    "agency",
    "marketing",
    "industries",
)
GENERIC_ENTITY_TOKENS = {
    "the",
    "sdn",
    "bhd",
    "co",
    "company",
    "enterprise",
    "trading",
    "supplies",
    "supply",
    "resources",
    "services",
    "service",
    "holdings",
    "restaurant",
    "cafe",
    "kitchen",
    "food",
    "distribution",
    "industries",
    "marketing",
}
BANK_TEXT_LINE_PATTERN = re.compile(
    r"^\s*(?P<date>(?:\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?)|(?:\d{1,2}\s+[A-Za-z]{3}(?:\s+\d{2,4})?))\s+(?P<rest>.+)$"
)
AMOUNT_TOKEN_PATTERN = re.compile(r"(?:RM\s*)?\(?(?:(?:\d{1,3}(?:,\d{3})+)|\d+)(?:\.\d{2})?\)?", re.I)
LABELED_PARTY_PATTERN = re.compile(
    r"(?:supplier|vendor|merchant|bill to|customer|company|issued by|from)\s*[:=-]\s*(.+)",
    re.I,
)
LABELED_DATE_PATTERN = re.compile(
    r"(?:invoice date|bill date|receipt date|statement date|date)\s*[:=-]\s*([^\n]+)",
    re.I,
)
LABELED_AMOUNT_PATTERN = re.compile(
    r"(?:total due|amount due|invoice total|net total|total payable|amount payable|grand total|total|net sales|total sales)\s*[:=-]?\s*((?:RM\s*)?\(?(?:(?:\d{1,3}(?:,\d{3})+)|\d+)(?:\.\d{2})?\)?)",
    re.I,
)
LABELED_DELIVERY_DATE_PATTERN = re.compile(
    r"(?:delivery date|delivered on)\s*[:=-]\s*([^\n]+)",
    re.I,
)
AMOUNT_VALUE_PATTERN = re.compile(r"(?<!\d)(?:\d{1,3}(?:,\d{3})+|\d+)\.\d{2}(?!\d)")
OCR_AMOUNT_VALUE_PATTERN = re.compile(r"(?<!\d)(?:\d{1,3}(?:,\d{3})+|\d+)(?:[.\s]\d{2})(?!\d)")
INVOICE_NUMBER_PATTERN = re.compile(
    r"(?:invoice\s*(?:no|number)?|bill\s*no|doc(?:ument)?\s*no|cash sale\s*/\s*invoice|no\.?)\s*[:#-]?\s*([A-Z0-9][A-Z0-9/-]{2,})",
    re.I,
)
INVOICE_DATE_PATTERN = re.compile(
    r"(?:invoice date|bill date|receipt date|statement date|delivery date|date)\s*[:=-]?\s*([0-9]{1,2}[/-][0-9]{1,2}[/-][0-9]{2,4}|[0-9]{1,2}\s*[A-Za-z]{3,9}\s*[0-9]{0,4}|[A-Za-z]{3,9}\s*[0-9]{1,2},?\s*[0-9]{0,4})",
    re.I,
)
HEADER_COMPANY_PATTERN = re.compile(
    r"(?m)^\s*([A-Z][A-Z0-9&'().,/ -]{5,}(?:SDN\.?\s*BHD\.?|BHD\.?|ENTERPRISE|TRADING|SERVICES|SUPPLIES|SUPPLY|RESOURCES|RESTAURANT|CAFE|KITCHEN).*)$"
)
DELIVER_TO_BLOCK_PATTERN = re.compile(
    r"deliver(?:y)?\s+to\s*:?\s*(.+?)(?=\n\s*(?:your p/?o|our d/?o|terms|date|page|fax|tel|driver|lorry|invoice|no\.?)\b|\Z)",
    re.I | re.S,
)
CUSTOMER_BLOCK_PATTERN = re.compile(
    r"(?:bill to|invoice to|customer)\s*:?\s*(.+?)(?=\n\s*(?:deliver(?:y)? to|your p/?o|our d/?o|terms|date|page|fax|tel|invoice|no\.?)\b|\Z)",
    re.I | re.S,
)
DELIVERY_ORDER_PATTERNS = (
    re.compile(r"\bdelivery\s+order\b", re.I),
    re.compile(r"\bdelivery\s+note\b", re.I),
)
INVOICE_TOTAL_HINTS = (
    "grand total",
    "invoice total",
    "amount due",
    "total due",
    "total payable",
    "balance due",
    "balance",
    "balance(rm)",
    "balance (rm)",
    "net total",
    "total amount",
    "total/jumlah",
    "jumlah",
)
NON_AMOUNT_CONTEXT_HINTS = (
    "invoice no",
    "invoice :",
    "invoice:",
    "date :",
    "date:",
    "page :",
    "page:",
    "phone",
    "fax",
    "h/p",
    "tin no",
    "company no",
    "customer po",
    "salesperson",
    "printed by",
    "acct no",
    "account no",
    "a/c no",
)


def _complete_text_confidence(
    *,
    detected_type: str,
    target: str,
    date: str,
    party: str,
    reference: str,
    amount: float,
    suggested_gl_account: str,
) -> float:
    score = 0.0
    score += 0.22 if date else 0.0
    score += 0.23 if amount > 0 else 0.0
    score += 0.2 if party else 0.0
    score += 0.12 if reference else 0.0
    score += 0.13 if detected_type != "Unknown" else 0.0
    score += 0.1 if target == "WP2" or suggested_gl_account else 0.0
    return min(score, 0.96)


def _collapse_block(block: str) -> str:
    parts = [normalize_whitespace(line) for line in block.splitlines() if normalize_whitespace(line)]
    return " ".join(parts[:4]).strip()


def _extract_labeled_value(pattern: re.Pattern[str], text: str) -> str:
    match = pattern.search(text)
    if not match:
        return ""
    return match.group(1).strip(" :-")


def _extract_invoice_header_company(preview: str) -> str:
    header_region = preview[: max(preview.upper().find("INVOICE"), 0) or 1600]
    for match in HEADER_COMPANY_PATTERN.findall(header_region):
        candidate = normalize_whitespace(match)
        if candidate:
            return candidate
    return ""


def _extract_invoice_reference(preview: str, filename: str) -> str:
    code_pattern = r"([A-Z]{1,5}(?:[-/]?)\d{4,}(?:[/-][A-Z0-9]+)*)"
    filename_upper = filename.upper()
    filename_code_match = re.search(r"\b([A-Z]{1,8}(?:[-/]?)\d{4,}(?:[/-][A-Z0-9]+)*)\b", filename_upper)
    if not filename_code_match:
        filename_code_match = re.search(r"\b(\d{6,}(?:[/-][A-Z0-9]+)*)\b", filename_upper)
    filename_candidate = clean_reference(filename_code_match.group(1)) if filename_code_match else extract_reference("", filename)
    strong_patterns = (
        re.compile(rf"invoice\s*(?:no|number)?\s*[:#-]?\s*{code_pattern}", re.I),
        re.compile(rf"(?:invoice\s*(?:no|number)?|bill\s*no|doc(?:ument)?\s*no)\s*[:#-]?\s*{code_pattern}", re.I),
        re.compile(rf"(?:our\s*d/?o\s*no|customer\s*p/?o\s*:?|ref(?:erence)?)\s*[:#-]?\s*{code_pattern}", re.I),
        re.compile(rf"(?m)^\s*no\.?\s*[:#-]?\s*{code_pattern}\s*$", re.I),
    )
    for pattern in strong_patterns:
        match = pattern.search(preview)
        if match:
            return clean_reference(match.group(1))
    if filename_candidate and any(character.isdigit() for character in filename_candidate):
        return filename_candidate
    generic_pattern = re.compile(rf"\b{code_pattern}\b")
    generic_match = generic_pattern.search(preview)
    if generic_match:
        return clean_reference(generic_match.group(1))
    return extract_reference(preview[:4000], filename)


def _extract_company_candidates(preview: str) -> list[str]:
    candidates: list[str] = []
    lines = [normalize_whitespace(line) for line in preview.splitlines() if normalize_whitespace(line)]
    combined_lines: list[str] = []
    for index, line in enumerate(lines):
        combined_lines.append(line)
        if index + 1 < len(lines):
            combined_lines.append(normalize_whitespace(f"{line} {lines[index + 1]}"))
        if index + 2 < len(lines):
            combined_lines.append(normalize_whitespace(f"{line} {lines[index + 1]} {lines[index + 2]}"))

    for line in combined_lines:
        candidate = normalize_whitespace(line)
        lowered = normalize(candidate)
        if not candidate or len(candidate) < 6:
            continue
        if any(token in lowered for token in ("phone", "tel", "fax", "email", "company no", "tin no", "sst no", "reg no")):
            continue
        if any(token in lowered for token in COMPANY_SUFFIXES):
            candidates.append(candidate)
    return candidates


def _entity_strings_match(left: str, right: str) -> bool:
    normalized_left = normalize(left)
    normalized_right = normalize(right)
    if not normalized_left or not normalized_right:
        return False
    left_tokens = {
        token for token in normalized_left.split()
        if len(token) >= 4 and token not in GENERIC_ENTITY_TOKENS
    }
    right_tokens = {
        token for token in normalized_right.split()
        if len(token) >= 4 and token not in GENERIC_ENTITY_TOKENS
    }
    if left_tokens and right_tokens and left_tokens.intersection(right_tokens):
        return True
    return normalized_left in normalized_right or normalized_right in normalized_left


def _looks_like_delivery_order(preview: str, filename: str) -> bool:
    source = f"{filename}\n{preview[:4000]}"
    if any(pattern.search(source) for pattern in DELIVERY_ORDER_PATTERNS):
        return True

    normalized_filename = normalize(filename)
    normalized_preview = normalize(preview[:1200])
    return bool(re.search(r"(^|[^a-z])do([^a-z]|$)", normalized_filename) and "invoice" not in normalized_preview)


def _score_party_candidate(candidate: str) -> int:
    cleaned = clean_party_value(candidate)
    if not cleaned:
        return -10_000

    lowered = normalize(cleaned)
    tokens = [token.strip(".,()") for token in cleaned.split() if token.strip(".,()")]
    normalized_tokens = [token.casefold() for token in tokens]
    suspicious_body_tokens = (
        "invoice no",
        "customer po",
        "deliver to",
        "bill to",
        "date",
        "page",
        "salesperson",
        "terms",
        "attn",
        "all cheque",
        "all cheques",
        "paymen",
        "payable",
        "authorised signature",
        "past due",
        "eligible",
        "experian",
        "information services",
        "payment accepted",
    )
    score = len(cleaned)
    if "sdn bhd" in lowered:
        score += 120
    elif any(token in lowered for token in COMPANY_SUFFIXES):
        score += 60
    if any(character.isdigit() for character in cleaned):
        score -= 25
    if any(token in lowered for token in suspicious_body_tokens):
        score -= 120
    if ":" in candidate:
        score -= 25
    if cleaned.endswith((" SDN", " T", " TR", " TRA", " TRAD", " TRADI", " TRADIN")):
        score -= 40
    suffix_index = -1
    for index in range(len(normalized_tokens) - 1):
        if normalized_tokens[index] == "sdn" and normalized_tokens[index + 1] == "bhd":
            suffix_index = index
            break
    if suffix_index < 0:
        for index, token in enumerate(normalized_tokens):
            if token in COMPANY_SUFFIXES:
                suffix_index = index
                break
    if suffix_index > 0:
        pre_suffix_tokens = [token for token in tokens[:suffix_index] if any(character.isalpha() for character in token)]
        score += min(len(pre_suffix_tokens), 4) * 12
        if pre_suffix_tokens and len(pre_suffix_tokens[0]) <= 2:
            score -= 30
    elif tokens and len(tokens[0]) <= 2:
        score -= 35
    return score


def _invoice_amount_candidates_from_line(line: str) -> list[float]:
    return [parse_amount(match.group(0)) for match in OCR_AMOUNT_VALUE_PATTERN.finditer(line)]


def _invoice_amount_context_score(line: str, previous_line: str = "", next_line: str = "") -> int:
    context = " ".join(part for part in (previous_line, line, next_line) if part)
    lowered = normalize(context)
    score = 0

    if any(token in lowered for token in ("grand total", "invoice total", "amount due", "total due", "total payable", "net total")):
        score += 220
    if "total amount" in lowered or ("otal" in lowered and "aout" in lowered):
        score += 210
    elif re.search(r"\btotal\b", lowered) or " otal " in f" {lowered} ":
        score += 120
    if "subtotal" in lowered or "subtot" in lowered:
        score += 45
    if any(token in lowered for token in ("qty", "quantity", "u price", "unit price", "description", "sku")):
        score -= 80
    if any(token in lowered for token in ("paid", "payment accepted", "payment information", "balance", "remaining balance", "pending", "bank:", "bank ")):
        score -= 110
    if any(token in lowered for token in ("reference", "invoice no", "date", "page", "phone", "fax")):
        score -= 50
    return score


def _extract_invoice_footer_amount(lines: list[str]) -> float:
    if not lines:
        return 0.0

    footer_lines = lines[-30:]
    best_amount = 0.0
    for line in footer_lines:
        lowered = normalize(line)
        if any(token in lowered for token in NON_AMOUNT_CONTEXT_HINTS):
            continue
        for amount in _invoice_amount_candidates_from_line(line):
            if amount > best_amount:
                best_amount = amount
    return best_amount


def _preview_contains_decimal_amount(preview: str, amount: float) -> bool:
    if amount <= 0:
        return False
    for match in OCR_AMOUNT_VALUE_PATTERN.finditer(preview):
        if abs(parse_amount(match.group(0)) - amount) < 0.01:
            return True
    return False


def _extract_invoice_total_amount(preview: str, *, delivery_order_like: bool = False) -> float:
    lines = [normalize_whitespace(line) for line in preview.splitlines() if normalize_whitespace(line)]
    best_scored_amount = 0.0
    best_scored_value = -10_000

    for index, line in enumerate(lines):
        previous_line = lines[index - 1] if index > 0 else ""
        next_line = lines[index + 1] if index + 1 < len(lines) else ""
        candidates = _invoice_amount_candidates_from_line(line)
        if not candidates and next_line:
            candidates = _invoice_amount_candidates_from_line(next_line)
            if candidates:
                line_for_score = f"{line} {next_line}"
            else:
                line_for_score = line
        else:
            line_for_score = line
        if not candidates:
            continue
        context_score = _invoice_amount_context_score(line_for_score, previous_line, next_line)
        for amount in candidates:
            amount_score = context_score
            if amount >= 100:
                amount_score += 12
            elif amount >= 10:
                amount_score += 4
            if amount_score > best_scored_value or (amount_score == best_scored_value and amount > best_scored_amount):
                best_scored_amount = amount
                best_scored_value = amount_score

    if best_scored_amount > 0 and best_scored_value >= 110:
        return best_scored_amount

    footer_start = max(0, len(lines) // 2)
    for index in range(len(lines) - 1, footer_start - 1, -1):
        lowered = normalize(lines[index])
        if not any(keyword in lowered for keyword in INVOICE_TOTAL_HINTS):
            continue
        matches = _invoice_amount_candidates_from_line(lines[index])
        if matches:
            amount = matches[-1]
            if amount > 0:
                return amount
        if index + 1 < len(lines):
            next_matches = _invoice_amount_candidates_from_line(lines[index + 1])
            if next_matches:
                amount = next_matches[-1]
                if amount > 0:
                    return amount

    for line in reversed(lines):
        lowered = normalize(line)
        if lowered.startswith("total ") or lowered == "total":
            if delivery_order_like and "qty" in lowered:
                continue
            if "qty" in lowered or "quantity" in lowered:
                continue
            matches = _invoice_amount_candidates_from_line(line)
            if matches:
                amount = matches[-1]
                if amount > 0:
                    return amount

    for index, line in enumerate(lines):
        lowered = normalize(line)
        if lowered != "total":
            continue
        if index + 1 >= len(lines):
            continue
        next_line = lines[index + 1]
        next_matches = _invoice_amount_candidates_from_line(next_line)
        if next_matches:
            amount = next_matches[-1]
            if amount > 0:
                return amount

    if delivery_order_like:
        return 0.0

    footer_amount = _extract_invoice_footer_amount(lines)
    if footer_amount > 0:
        return footer_amount

    fallback_amount = extract_document_amount(preview)
    if fallback_amount <= 1.0:
        return 0.0
    if not _preview_contains_decimal_amount(preview, fallback_amount):
        return 0.0
    return fallback_amount


def _filter_generic_invoice_warnings(
    warnings: list[str],
    *,
    has_strong_fields: bool,
    detection_type: str,
) -> list[str]:
    if not warnings:
        return warnings
    if not has_strong_fields:
        return warnings
    if detection_type not in {"Purchase Invoice", "Sales Invoice"}:
        return warnings

    suppressed_prefixes = (
        "Only one strong clue found.",
        "Also looked like ",
        "Bank statement was inferred from filename/text clues.",
        "Invoice type was inferred from filename/text clues.",
        "Sales invoice type was inferred from filename/text clues.",
    )
    return [
        warning for warning in warnings
        if not any(warning.startswith(prefix) for prefix in suppressed_prefixes)
    ]


def _looks_like_structured_invoice(preview: str, filename: str) -> bool:
    normalized_preview = normalize(preview[:6000])
    normalized_filename = normalize(filename)
    if "invoice" not in normalized_preview and "invoice" not in normalized_filename and "tax invoice" not in normalized_preview:
        return False

    invoice_signals = (
        "deliver to",
        "bill to",
        "ship to",
        "your p/o no",
        "our d/o no",
        "terms",
        "payment term",
        "subtotal",
        "total qty",
        "description uom",
        "invoice date",
        "amount due",
        "total due",
    )
    signal_hits = sum(1 for token in invoice_signals if token in normalized_preview)
    return signal_hits >= 2


def _party_from_filename(filename: str) -> str:
    cleaned = re.sub(r"\.[^.]+$", "", filename)
    cleaned = re.sub(r"[_-]+", " ", cleaned)
    cleaned = re.sub(r"\b(?:inv|invoice|receipt|statement|tax|official|row)\b", " ", cleaned, flags=re.I)
    tokens = [token for token in cleaned.split() if token and not token.isdigit()]
    return " ".join(tokens[:6]).strip()


def extract_structured_document_fields(filename: str, text: str, detection: dict[str, Any]) -> dict[str, Any]:
    preview = text[:5000]
    date = normalize_date(_extract_labeled_value(LABELED_DATE_PATTERN, preview)) or normalize_date(preview)
    if not date:
        date = normalize_date(_extract_labeled_value(LABELED_DELIVERY_DATE_PATTERN, preview))
    reference = extract_reference(preview[:3000], filename)
    party = extract_document_party(filename, preview)

    labeled_amount = _extract_labeled_value(LABELED_AMOUNT_PATTERN, preview)
    amount = parse_amount(labeled_amount) if labeled_amount else 0.0
    if amount <= 0:
        amount = extract_document_amount(f"{filename}\n{preview}")

    evidence: list[str] = list(detection.get("evidence", []))
    if labeled_amount and amount > 0:
        evidence.append("Amount came from a labeled total/amount field in the document.")
    if party:
        evidence.append("Party was taken from the document header or filename.")
    if date:
        evidence.append("Date was detected from labeled date fields or readable document text.")

    return {
        "date": date,
        "reference": reference,
        "party": party,
        "amount": amount,
        "evidence": evidence,
        "fieldConfidence": {
            "date": 0.94 if date else 0.0,
            "reference": 0.9 if reference else 0.0,
            "party": 0.88 if party else 0.0,
            "amount": 0.95 if amount > 0 else 0.0,
            "suggestedGlAccount": 0.86 if detection.get("suggestedGlAccount") else 0.0,
        },
        "fieldEvidence": {
            "date": ["Detected from labeled date field or readable document text."] if date else [],
            "reference": ["Detected from invoice/reference/account markers or strong fallback pattern."] if reference else [],
            "party": ["Detected from company/vendor lines near the document header."] if party else [],
            "amount": ["Detected from labeled total/amount fields."] if amount > 0 else [],
        },
    }


def extract_invoice_text_item(
    filename: str,
    content_type: str,
    data: bytes,
    text: str,
    upload_lane: str | None = None,
    client_entity_name: str | None = None,
) -> list[dict[str, Any]]:
    preview = text[:12000]
    normalized_preview = normalize(preview)
    normalized_filename = normalize(filename)
    lane = normalize_upload_lane(upload_lane)

    invoice_like = any(
        token in normalized_preview or token in normalized_filename
        for token in ("invoice", "tax invoice", "cash sale", "bill to", "deliver to", "invoice date", "amount due")
    )
    if not invoice_like:
        return []

    detection = detect_type(f"{filename}\n{preview[:6000]}", lane)
    invoice_role, role_evidence, role_warnings = infer_invoice_role_from_text(
        filename=filename,
        text=preview,
        client_entity_name=client_entity_name or "",
    )

    issuer = clean_party_value(_extract_invoice_header_company(preview))

    customer_block_match = CUSTOMER_BLOCK_PATTERN.search(preview[:4000])
    deliver_to_match = DELIVER_TO_BLOCK_PATTERN.search(preview[:4000])
    customer_block = _collapse_block(customer_block_match.group(1)) if customer_block_match else ""
    deliver_to_block = _collapse_block(deliver_to_match.group(1)) if deliver_to_match else ""
    strong_invoice_layout = _looks_like_structured_invoice(preview, filename)
    company_candidates = [
        candidate
        for candidate in (clean_party_value(value) for value in _extract_company_candidates(preview[:5000]))
        if candidate
    ]
    billed_party = customer_block or deliver_to_block
    supplier_pool = [
        candidate
        for candidate in company_candidates
        if not billed_party or not _entity_strings_match(candidate, billed_party)
    ]
    strong_supplier_pool = [
        (index, candidate)
        for index, candidate in enumerate(supplier_pool)
        if (
            "sdn bhd" in normalize(candidate)
            or any(token in normalize(candidate) for token in COMPANY_SUFFIXES[2:])
        )
        and not normalize(candidate).startswith(("for ", "by for ", "acknowle by "))
    ]
    supplier_candidate = (
        max(strong_supplier_pool, key=lambda item: (_score_party_candidate(item[1]), -item[0]))[1]
        if strong_supplier_pool
        else max(supplier_pool, key=_score_party_candidate, default="")
    )
    client_matches_issuer = bool(client_entity_name and issuer and _client_match_strength(issuer, client_entity_name))
    distinct_two_party_invoice = bool(billed_party and supplier_candidate and not _entity_strings_match(supplier_candidate, billed_party))

    if invoice_role == "purchase" or (lane == "purchases" and detection["detectedType"] in {"Unknown", "Sales Invoice", "Purchase Invoice"}):
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
    elif invoice_role == "sales" and (client_matches_issuer or not distinct_two_party_invoice):
        detection["detectedType"] = "Sales Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "4100 - Sales Revenue"
    elif invoice_role == "sales":
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
        detection["warnings"] = [
            *detection.get("warnings", []),
            "Sales-side classification was downgraded because the invoice showed a separate billed party and supplier.",
        ]
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Distinct supplier and billed-party fields made the invoice look purchase-side.",
        ]
    elif lane == "sales" and detection["detectedType"] in {"Unknown", "Sales Invoice"} and client_matches_issuer and not distinct_two_party_invoice:
        detection["detectedType"] = "Sales Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "4100 - Sales Revenue"
    elif lane == "sales" and distinct_two_party_invoice:
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Distinct supplier and billed-party fields blocked a sales-lane override.",
        ]
    elif lane == "purchases":
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
    elif strong_invoice_layout and billed_party and supplier_candidate and detection["detectedType"] != "Sales Invoice":
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Invoice showed a billed customer and a different supplier company line, so it was treated as a purchase invoice.",
        ]
    elif strong_invoice_layout and detection["detectedType"] != "Sales Invoice":
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Structured invoice layout was strong enough to treat the document as a purchase invoice by default.",
        ]
    elif detection["detectedType"] == "Bank Statement" and strong_invoice_layout:
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Structured invoice layout overrode a weaker bank-style keyword match.",
        ]
        detection["warnings"] = [
            warning
            for warning in detection.get("warnings", [])
            if warning != "Also looked like Purchase Invoice."
        ]
        detection["warnings"].append("Invoice layout overrode a weaker bank-style clue. Review sales versus purchase only if the file was uploaded to the wrong client session.")

    date_match = INVOICE_DATE_PATTERN.search(preview)
    labeled_date = normalize_date(date_match.group(1)) if date_match else ""
    if not labeled_date:
        labeled_date = normalize_date(_extract_labeled_value(LABELED_DELIVERY_DATE_PATTERN, preview))
    if not labeled_date:
        labeled_date = normalize_date(_extract_labeled_value(LABELED_DATE_PATTERN, preview)) or normalize_date(preview)

    reference = _extract_invoice_reference(preview, filename)

    delivery_order_like = _looks_like_delivery_order(preview, filename)
    amount = _extract_invoice_total_amount(preview, delivery_order_like=delivery_order_like)
    if amount <= 0 and not delivery_order_like:
        amount = extract_document_amount(f"{filename}\n{preview[:4000]}")
    if delivery_order_like and amount <= 0:
        detection["detectedType"] = "Receipt" if detection.get("detectedType") == "Receipt" else "Unknown"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = ""
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Delivery-note style layout was separated from invoice posting because no reliable posting total was found.",
        ]

    fallback_party = clean_party_value(extract_document_party(filename, preview))
    party = supplier_candidate or issuer or fallback_party
    if not party and customer_block:
        party = customer_block

    if detection["detectedType"] == "Bank Statement" and reference and amount > 0:
        detection["detectedType"] = "Purchase Invoice" if invoice_role != "sales" else "Sales Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials" if detection["detectedType"] == "Purchase Invoice" else "4100 - Sales Revenue"
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Invoice parser overrode a weaker bank-style keyword match because invoice fields were present.",
        ]
        detection["warnings"] = [
            warning
            for warning in detection.get("warnings", [])
            if warning != "Bank statement was inferred from filename/text clues. Review before posting."
        ]

    evidence = [*detection.get("evidence", []), *role_evidence]
    if issuer:
        evidence.append("Supplier / issuer name came from the invoice header.")
    if customer_block:
        evidence.append("Customer / Bill To block was read from the invoice body.")
    if deliver_to_block:
        evidence.append("Deliver To block was read from the invoice body.")
    if reference:
        evidence.append("Reference came from the visible invoice/document number.")
    if labeled_date:
        evidence.append("Date came from the visible invoice/date field.")
    if amount > 0:
        evidence.append("Amount came from invoice total / amount fields in the readable PDF text.")

    warnings = [*detection.get("warnings", []), *role_warnings]
    if not labeled_date:
        warnings.append("Date was not detected.")
    if amount <= 0:
        warnings.append("Amount was not detected.")
    if delivery_order_like and amount <= 0:
        warnings.append("Delivery-order style document did not expose a reliable posting amount.")

    warnings = _filter_generic_invoice_warnings(
        warnings,
        has_strong_fields=bool(labeled_date and reference and amount > 0 and party),
        detection_type=detection["detectedType"],
    )

    confidence = _complete_text_confidence(
        detected_type=detection["detectedType"],
        target=detection["target"],
        date=labeled_date,
        party=party,
        reference=reference,
        amount=amount,
        suggested_gl_account=detection["suggestedGlAccount"],
    )
    confidence = max(confidence, 0.84 if detection["detectedType"] != "Unknown" and amount > 0 else confidence)

    return [
        build_item(
            file_name=filename,
            file_type=content_type or "application/pdf",
            file_size=len(data),
            client_entity_name=client_entity_name,
            detected_type=detection["detectedType"],
            target=detection["target"],
            date=labeled_date,
            reference=reference,
            party=party,
            amount=amount,
            suggested_gl_account=detection["suggestedGlAccount"],
            notes="Extracted from readable invoice PDF text using the supplier-invoice parser.",
            evidence=evidence,
            warnings=[warning for warning in warnings if warning],
            raw_preview=preview[:240],
            overall_confidence_score=confidence,
            extraction_method="pdf-invoice-text-parser",
            field_confidence={
                "date": 0.95 if labeled_date else 0.0,
                "reference": 0.94 if reference else 0.0,
                "party": 0.92 if party else 0.0,
                "amount": 0.96 if amount > 0 else 0.0,
                "suggestedGlAccount": 0.9 if detection["suggestedGlAccount"] else 0.0,
            },
            field_evidence={
                "date": ["Detected from visible invoice/date field."] if labeled_date else [],
                "reference": ["Detected from visible invoice/document number."] if reference else [],
                "party": ["Detected from supplier header or invoice party blocks."] if party else [],
                "amount": ["Detected from invoice total / amount field."] if amount > 0 else [],
            },
        )
    ]


def _client_match_strength(text: str, client_entity_name: str) -> bool:
    normalized_client = normalize(client_entity_name)
    normalized_text = normalize(text)
    if not normalized_client or not normalized_text:
        return False
    client_tokens = [token for token in normalized_client.split() if len(token) > 2]
    if not client_tokens:
        return False
    distinctive_tokens = [
        token for token in client_tokens
        if token not in GENERIC_ENTITY_TOKENS and len(token) >= 4
    ]
    if distinctive_tokens:
        return any(token in normalized_text for token in distinctive_tokens)
    return sum(1 for token in client_tokens if token in normalized_text) >= max(2, min(3, len(client_tokens)))


def infer_invoice_role_from_text(
    *,
    filename: str,
    text: str,
    client_entity_name: str = "",
) -> tuple[str, list[str], list[str]]:
    preview = text[:6000]
    normalized_preview = normalize(preview)
    normalized_filename = normalize(filename)
    evidence: list[str] = []
    warnings: list[str] = []

    if "invoice" not in normalized_preview and "invoice" not in normalized_filename and "tax invoice" not in normalized_preview:
        return "unknown", evidence, warnings

    client_matches_bill_to = False
    client_matches_deliver_to = False
    client_matches_issuer = False
    client_matches_body = False
    customer_block_match = CUSTOMER_BLOCK_PATTERN.search(preview[:4000])
    deliver_to_match = DELIVER_TO_BLOCK_PATTERN.search(preview[:4000])
    customer_block = _collapse_block(customer_block_match.group(1)) if customer_block_match else ""
    deliver_to_block = _collapse_block(deliver_to_match.group(1)) if deliver_to_match else ""
    issuer_header = _extract_invoice_header_company(preview)
    company_candidates = _extract_company_candidates(preview[:5000])

    if client_entity_name:
        invoice_body = preview[max(preview.upper().find("INVOICE"), 0):4000]

        client_matches_bill_to = bool(customer_block and _client_match_strength(customer_block, client_entity_name))
        client_matches_deliver_to = bool(deliver_to_block and _client_match_strength(deliver_to_block, client_entity_name))
        client_matches_issuer = bool(issuer_header and _client_match_strength(issuer_header, client_entity_name))
        client_matches_body = _client_match_strength(invoice_body, client_entity_name)

        if client_matches_bill_to:
            evidence.append("Client entity matched a Bill To / Customer field.")
        if client_matches_deliver_to:
            evidence.append("Client entity matched a Deliver To field.")
        if client_matches_issuer:
            evidence.append("Client entity matched the issuer/header area.")
        if client_matches_body and not client_matches_issuer:
            evidence.append("Client entity was present in the invoice body but not the supplier header.")

    if client_matches_bill_to or client_matches_deliver_to:
        return "purchase", evidence, warnings
    if client_matches_body and not client_matches_issuer:
        return "purchase", evidence, warnings
    if client_matches_issuer:
        return "sales", evidence, warnings

    billed_party = customer_block or deliver_to_block
    if billed_party:
        distinct_supplier = next(
            (
                candidate for candidate in company_candidates
                if not _entity_strings_match(candidate, billed_party)
            ),
            "",
        )
        if distinct_supplier:
            evidence.append("Invoice showed a billed customer and a different supplier company line.")
            return "purchase", evidence, warnings

    content = normalize(f"{filename} {preview[:2500]}")
    if any(token in content for token in ("supplier", "vendor", "bill from", "purchase invoice", "tax invoice")):
        evidence.append("Detected supplier-side invoice cues.")
        return "purchase", evidence, warnings
    if client_matches_issuer and any(token in content for token in ("sales invoice", "invoice to", "customer invoice")):
        evidence.append("Detected sales-side invoice cues.")
        return "sales", evidence, warnings

    return "unknown", evidence, warnings


def extract_spreadsheet_items(
    filename: str,
    content_type: str,
    data: bytes,
    upload_lane: str | None = None,
    client_entity_name: str | None = None,
) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return []
    items: list[dict[str, Any]] = []
    uploaded_at = current_timestamp()

    for sheet in workbook.worksheets:
        rows = [
            [cell_to_text(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
            if any(cell_to_text(value) for value in row)
        ]
        header_index = find_header_index(rows)
        if header_index is None:
            continue

        raw_headers = rows[header_index]
        headers = [normalize(header) for header in raw_headers]
        sample_rows = [" ".join(row) for row in rows[header_index + 1 : header_index + 13] if any(row)]
        context_detection = classify_tabular_context(
            filename=f"{filename} {sheet.title}",
            headers=headers,
            sample_rows=sample_rows,
            row_count=max(0, len(rows) - header_index - 1),
            upload_lane=upload_lane,
        )
        for offset, row in enumerate(rows[header_index + 1 : header_index + 101]):
            if not any(row):
                continue
            actual_row_index = header_index + 1 + offset
            source_row_number = actual_row_index + 1
            context_start = max(header_index + 1, actual_row_index - 2)
            context_end = min(len(rows), actual_row_index + 3)
            context_rows = [
                {
                    "row_number": row_number + 1,
                    "cells": context_row,
                }
                for row_number, context_row in enumerate(rows[context_start:context_end], start=context_start)
                if any(context_row)
            ]
            item = item_from_cells(
                filename=f"{filename} {sheet.title} row {offset + 1}",
                content_type=content_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                file_size=len(data),
                uploaded_at=uploaded_at,
                headers=headers,
                raw_headers=raw_headers,
                cells=row,
                row_text=" ".join(row),
                extraction_method="spreadsheet",
                context_detection=context_detection,
                upload_lane=upload_lane,
                client_entity_name=client_entity_name,
                source_file_name=f"{filename} {sheet.title}",
                header_row_number=header_index + 1,
                source_row_number=source_row_number,
                raw_context_rows=context_rows,
            )
            if item:
                items.append(item)

    return items


def extract_delimited_items(
    filename: str,
    content_type: str,
    text: str,
    upload_lane: str | None = None,
    client_entity_name: str | None = None,
) -> list[dict[str, Any]]:
    rows = split_delimited_rows(text)
    if len(rows) < 2:
        return []

    header_index = find_header_index(rows)
    if header_index is None:
        return []

    raw_headers = rows[header_index]
    headers = [normalize(header) for header in raw_headers]
    sample_rows = [" ".join(row) for row in rows[header_index + 1 : header_index + 13] if any(cell.strip() for cell in row)]
    context_detection = classify_tabular_context(
        filename=filename,
        headers=headers,
        sample_rows=sample_rows,
        row_count=max(0, len(rows) - header_index - 1),
        upload_lane=upload_lane,
    )
    uploaded_at = current_timestamp()
    items: list[dict[str, Any]] = []
    for index, row in enumerate(rows[header_index + 1 : header_index + 101]):
        if not any(cell.strip() for cell in row):
            continue
        actual_row_index = header_index + 1 + index
        context_start = max(header_index + 1, actual_row_index - 2)
        context_end = min(len(rows), actual_row_index + 3)
        context_rows = [
            {
                "row_number": row_number + 1,
                "cells": context_row,
            }
            for row_number, context_row in enumerate(rows[context_start:context_end], start=context_start)
            if any(cell.strip() for cell in context_row)
        ]
        item = item_from_cells(
            filename=f"{filename} row {index + 1}",
            content_type=content_type or "text/csv",
            file_size=len(text.encode("utf-8")),
            uploaded_at=uploaded_at,
            headers=headers,
            raw_headers=raw_headers,
            cells=row,
            row_text=" ".join(row),
            extraction_method="delimited-text",
            context_detection=context_detection,
            upload_lane=upload_lane,
            client_entity_name=client_entity_name,
            source_file_name=filename,
            header_row_number=header_index + 1,
            source_row_number=actual_row_index + 1,
            raw_context_rows=context_rows,
        )
        if item:
            items.append(item)
    return items


# ── Public Bank card-terminal merchant statement parser ────────────────────────

_PB_MERCH_STMT_SIGNALS = ("merchant statement", "merchant id", "grand total for this statement")
_MERCH_ID_RE = re.compile(r"\bMERCHANT ID\s+(\d+)", re.I)
_MERCH_DATE_RE = re.compile(r"\bDATE\s+(\d{1,2})([A-Z]{3})(\d{2,4})\b", re.I)
_MERCH_GRAND_TOTAL_RE = re.compile(
    r"GRAND TOTAL FOR THIS STATEMENT\s+[\d,]+\s+[\d,]+\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d*)",
    re.I,
)
_MERCH_MONTH_NUM = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
# Daily batch row: DD MMM  BATCHNO  MERCH_LOC  MACHINE  TRANS  REJECT  GROSS  DISC  NET
_MERCH_BATCH_ROW_RE = re.compile(
    r"(\d{1,2})\s*([A-Z]{3})\s+"       # day  month
    r"(\d{4,6})"                         # batch number
    r"(?:\s+\d+){4}\s+"                  # 4 integer columns: loc, machine, trans, reject
    r"([\d,]+\.\d{2})\s+"               # gross amount
    r"([\d,]+\.\d{2})\s+"               # disc (MDR) amount
    r"([\d,]+\.\d{2})",                  # net amount
    re.I,
)


def _pb_merch_statement_date(text: str) -> str:
    m = _MERCH_DATE_RE.search(text)
    if not m:
        return ""
    day, mon = m.group(1), m.group(2).lower()
    return f"{day} {mon.capitalize()}" if mon in _MERCH_MONTH_NUM else ""


def extract_merchant_statement_items(
    filename: str,
    text: str,
    *,
    content_type: str = "application/pdf",
    data: bytes = b"",
    client_entity_name: str | None = None,
) -> list[dict[str, Any]]:
    """Parse a Public Bank card-terminal merchant statement.

    Primary path: one WP1 item per daily batch at NET amount (reference = batch number)
    so each item can match a DEP-MERCHANT bank credit in WP2.  Also emits one monthly
    Merchant Discount Fee item per terminal for the MDR expense.

    Fallback (if daily rows cannot be parsed): emits monthly gross + MDR items from
    the Grand Total row, preserving the original two-item-per-terminal behaviour.
    """
    norm_full = normalize(text)
    if not all(sig in norm_full for sig in _PB_MERCH_STMT_SIGNALS):
        return []

    all_merch_id_pos = [(m.group(1), m.start()) for m in _MERCH_ID_RE.finditer(text)]
    grand_totals_raw = [
        (m.group(1), m.group(2), m.group(3), m.start())
        for m in _MERCH_GRAND_TOTAL_RE.finditer(text)
    ]
    if not all_merch_id_pos or not grand_totals_raw:
        return []

    statement_date = _pb_merch_statement_date(text)
    items: list[dict[str, Any]] = []

    # MERCHANT ID repeats on every PDF page; build ordered list of unique terminals
    # and span each terminal's section from its first page to the first page of the next terminal.
    unique_terminal_order: list[str] = []
    first_pos: dict[str, int] = {}
    for mid, mpos in all_merch_id_pos:
        if mid not in first_pos:
            first_pos[mid] = mpos
            unique_terminal_order.append(mid)

    for t_idx, merch_id in enumerate(unique_terminal_order):
        sec_start = first_pos[merch_id]
        sec_end = first_pos[unique_terminal_order[t_idx + 1]] if t_idx + 1 < len(unique_terminal_order) else len(text)
        section_text = text[sec_start:sec_end]

        grand_total = next(
            ((gs, ds, ns) for gs, ds, ns, gtpos in grand_totals_raw if sec_start <= gtpos < sec_end),
            None,
        )

        # Per-day batch rows: one WP1 item per batch at NET amount (matches bank DEP-MERCHANT credit)
        batch_count = 0
        total_disc_from_rows = 0.0
        for m in _MERCH_BATCH_ROW_RE.finditer(section_text):
            day, mon = m.group(1), m.group(2).lower()
            if mon not in _MERCH_MONTH_NUM:
                continue
            batch = m.group(3)
            gross_str, disc_str, net_str = m.group(4), m.group(5), m.group(6)
            gross = parse_amount(gross_str.replace(",", ""))
            disc = parse_amount(disc_str.replace(",", ""))
            net = parse_amount(net_str.replace(",", ""))
            if net <= 0:
                continue
            date_str = f"{int(day):02d} {mon.capitalize()}"
            items.append(build_item(
                file_name=filename,
                file_type=content_type,
                file_size=len(data),
                client_entity_name=client_entity_name,
                detected_type="Merchant Statement",
                target="WP1",
                date=date_str,
                party="Public Bank",
                reference=batch,
                amount=net,
                suggested_gl_account="4100 - Sales Revenue",
                notes=f"Card sales deposit — Merchant ID {merch_id}, Batch {batch}. Gross: RM {gross:,.2f}, MDR: RM {disc:,.2f}, Net: RM {net:,.2f}.",
                evidence=[f"Batch row — {date_str}, Batch {batch}, Gross: {gross_str}, Disc: {disc_str}, Net: {net_str}"],
                raw_preview=text[:240],
                overall_confidence_score=0.90,
                extraction_method="text-parser",
            ))
            total_disc_from_rows += disc
            batch_count += 1

        if batch_count > 0:
            # Monthly MDR expense item from grand total (preferred) or sum of daily rows
            if grand_total:
                gs, ds, ns = grand_total
                mdr_disc = parse_amount(ds.replace(",", ""))
                mdr_gross = parse_amount(gs.replace(",", ""))
                mdr_net = parse_amount(ns.replace(",", ""))
                mdr_note = f"Gross: RM {mdr_gross:,.2f}, MDR: RM {mdr_disc:,.2f}, Net: RM {mdr_net:,.2f}."
                mdr_evidence = [f"Grand Total — Gross: {gs}, Disc: {ds}, Net: {ns}"]
                mdr_confidence = 0.90
            else:
                mdr_disc = total_disc_from_rows
                mdr_note = "Summed from daily rows."
                mdr_evidence = ["MDR summed from daily batch rows (no Grand Total row found)."]
                mdr_confidence = 0.85
            if mdr_disc > 0:
                items.append(build_item(
                    file_name=filename,
                    file_type=content_type,
                    file_size=len(data),
                    client_entity_name=client_entity_name,
                    detected_type="Merchant Discount Fee",
                    target="WP1",
                    date=statement_date,
                    party="Public Bank",
                    reference=f"MERCH-{merch_id}",
                    amount=mdr_disc,
                    suggested_gl_account="6370 - Bank Charges & Fees",
                    notes=f"Monthly MDR fee — Merchant ID {merch_id}. {mdr_note}",
                    evidence=mdr_evidence,
                    raw_preview=text[:240],
                    overall_confidence_score=mdr_confidence,
                    extraction_method="text-parser",
                ))

        else:
            # Fallback: daily rows not parsed — monthly gross + MDR from grand total
            if not grand_total:
                continue
            gs, ds, ns = grand_total
            gross = parse_amount(gs.replace(",", ""))
            disc = parse_amount(ds.replace(",", ""))
            net = parse_amount(ns.replace(",", ""))
            net_note = f"Net deposit: RM {net:,.2f}." if net > 0 else ""
            if gross > 0:
                items.append(build_item(
                    file_name=filename,
                    file_type=content_type,
                    file_size=len(data),
                    client_entity_name=client_entity_name,
                    detected_type="Merchant Statement",
                    target="WP1",
                    date=statement_date,
                    party="Public Bank",
                    reference=f"MERCH-{merch_id}",
                    amount=gross,
                    suggested_gl_account="4100 - Sales Revenue",
                    notes=f"Gross card sales (Merchant ID {merch_id}). {net_note}".strip(),
                    evidence=[f"Grand Total — Gross: {gs}, Disc: {ds}, Net: {ns}"],
                    raw_preview=text[:240],
                    overall_confidence_score=0.90,
                    extraction_method="text-parser",
                ))
            if disc > 0:
                items.append(build_item(
                    file_name=filename,
                    file_type=content_type,
                    file_size=len(data),
                    client_entity_name=client_entity_name,
                    detected_type="Merchant Discount Fee",
                    target="WP1",
                    date=statement_date,
                    party="Public Bank",
                    reference=f"MERCH-{merch_id}",
                    amount=disc,
                    suggested_gl_account="6370 - Bank Charges & Fees",
                    notes=f"MDR card processing fee (Merchant ID {merch_id}). {net_note}".strip(),
                    evidence=[f"Grand Total — Disc AMT: {ds}"],
                    raw_preview=text[:240],
                    overall_confidence_score=0.90,
                    extraction_method="text-parser",
                ))

    return items


def extract_text_item(
    filename: str,
    content_type: str,
    data: bytes,
    upload_lane: str | None = None,
    client_entity_name: str | None = None,
) -> list[dict[str, Any]]:
    text = decode_text(data)
    bank_rows = extract_bank_statement_rows_from_text(
        filename=filename,
        content_type=content_type or "text/plain",
        data=data,
        text=text,
        client_entity_name=client_entity_name,
    )
    if bank_rows:
        return bank_rows

    merchant_items = extract_merchant_statement_items(
        filename, text,
        content_type=content_type or "application/pdf",
        data=data,
        client_entity_name=client_entity_name,
    )
    if merchant_items:
        return merchant_items

    invoice_items = extract_invoice_text_item(
        filename,
        content_type or "text/plain",
        data,
        text,
        upload_lane=upload_lane,
        client_entity_name=client_entity_name,
    )
    if invoice_items:
        return invoice_items

    lane = normalize_upload_lane(upload_lane)
    detection = detect_type(f"{filename} {text[:3000]}", lane)
    revenue_detection = None

    _content_preview = normalize(f"{filename} {text[:3000]}")
    _is_merchant_doc = any(
        token in _content_preview
        for token in ("merchant statement", "merchant payout", "merchant settlement", "merch stmt", "merch_stmt")
    )

    if lane in {"auto", "sales"} or _is_merchant_doc:
        revenue_detection = classify_revenue_context(
            filename=filename,
            raw_text=text[:3000],
            sample_rows=text.splitlines()[:12],
        )
    if revenue_detection and (lane != "purchases" or _is_merchant_doc):
        detection = revenue_detection

    invoice_role, role_evidence, role_warnings = infer_invoice_role_from_text(
        filename=filename,
        text=text,
        client_entity_name=client_entity_name or "",
    )
    normalized_text = normalize(text[:4000])
    invoice_like_purchase_clues = any(
        token in normalized_text
        for token in ("invoice", "tax invoice", "invoice date", "amount due", "total qty", "deliver to", "bill to")
    )
    if invoice_role == "purchase" and not _is_merchant_doc:
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = detection.get("suggestedGlAccount") or "5020 - Direct Materials"
    elif invoice_role == "sales":
        detection["detectedType"] = "Sales Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = detection.get("suggestedGlAccount") or "4100 - Sales Revenue"
    elif detection["detectedType"] in {"Sales Summary", "Sales Invoice", "Unknown"} and invoice_like_purchase_clues and not _is_merchant_doc:
        detection["detectedType"] = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = detection.get("suggestedGlAccount") or "5020 - Direct Materials"
        detection["evidence"] = [
            *detection.get("evidence", []),
            "Invoice-style OCR cues blocked a revenue-style fallback classification.",
        ]
        detection["warnings"] = [
            *detection.get("warnings", []),
            "Invoice OCR was weak, so the document was kept purchase-side for review instead of revenue-side auto posting.",
        ]
    if role_evidence:
        detection["evidence"] = [*detection.get("evidence", []), *role_evidence]
    if role_warnings:
        detection["warnings"] = [*detection.get("warnings", []), *role_warnings]

    structured = extract_structured_document_fields(filename, text, detection)
    amount = structured["amount"]
    date = structured["date"]
    party = structured["party"]
    reference = structured["reference"]
    overall_confidence = max(
        detection["confidenceScore"],
        _complete_text_confidence(
            detected_type=detection["detectedType"],
            target=detection["target"],
            date=date,
            party=party,
            reference=reference,
            amount=amount,
            suggested_gl_account=detection["suggestedGlAccount"],
        ),
    )
    warnings = [
        *detection["warnings"],
        "" if date else "Date was not detected.",
        "" if amount > 0 else "Amount was not detected.",
    ]
    return [
        build_item(
            file_name=filename,
            file_type=content_type or "text/plain",
            file_size=len(data),
            client_entity_name=client_entity_name,
            detected_type=detection["detectedType"],
            target=detection["target"],
            date=date,
            party=party,
            reference=reference,
            amount=amount,
            money_in=amount if detection["target"] == "WP2" else 0,
            money_out=0,
            suggested_gl_account=detection["suggestedGlAccount"],
            notes="Extracted from readable text content.",
            evidence=structured["evidence"],
            warnings=[warning for warning in warnings if warning],
            raw_preview=text[:240],
            overall_confidence_score=overall_confidence,
            extraction_method="text-parser",
            gl_suggestions=detection.get("glSuggestions"),
            field_confidence=structured["fieldConfidence"],
            field_evidence=structured["fieldEvidence"],
        )
    ]


def classify_tabular_context(
    *,
    filename: str,
    headers: list[str],
    sample_rows: list[str],
    row_count: int,
    upload_lane: str | None = None,
) -> dict[str, Any] | None:
    lane = normalize_upload_lane(upload_lane)
    header_text = " ".join(headers)
    normalized_header_text = normalize(header_text)
    if lane == "bank" and looks_bank_like(f"{filename} {header_text} {' '.join(sample_rows[:6])}"):
        return {
            "detectedType": "Bank Statement",
            "confidenceScore": 0.96,
            "target": "WP2",
            "evidence": ["Upload lane was Bank, so the uploaded table was treated as bank-side support."],
            "warnings": [],
            "suggestedGlAccount": "",
            "glSuggestions": [],
        }
    looks_like_bank = bool(
        normalized_header_text
        and any(pattern in normalized_header_text for pattern in ("money in", "deposit", "credit", "money out", "withdrawal", "debit", "bank description", "balance"))
    )
    if looks_like_bank:
        return {
            "detectedType": "Bank Statement",
            "confidenceScore": 0.97,
            "target": "WP2",
            "evidence": ["Detected bank-style columns in uploaded file."],
            "warnings": [],
            "suggestedGlAccount": "",
            "glSuggestions": [],
        }

    looks_like_sales_summary = bool(
        ("sales" in normalized_header_text and any(pattern in normalized_header_text for pattern in ("category", "gross", "net", "discount", "qty", "quantity", "receipt", "transaction")))
        or any(pattern in normalize(filename) for pattern in ("sales by category", "daily sales", "sales report", "pos sales"))
    )
    if looks_like_sales_summary:
        if lane in {"auto", "sales"}:
            return classify_revenue_context(
                filename=filename,
                headers=headers,
                sample_rows=sample_rows,
                row_count=row_count,
            ) or {
                "detectedType": "Sales Summary",
                "confidenceScore": 0.95,
                "target": "WP1",
                "evidence": ["Detected sales-summary style columns in uploaded file."],
                "warnings": [],
                "suggestedGlAccount": "4100 - Sales Revenue",
                "glSuggestions": [],
            }
        if lane == "purchases":
            return {
                "detectedType": "Purchase Invoice",
                "confidenceScore": 0.64,
                "target": "WP1",
                "evidence": ["Sales-summary style columns were present, but the Purchases upload lane kept the table on the purchase side."],
                "warnings": ["Purchases upload lane overrode a sales-summary style guess. Review before posting."],
                "suggestedGlAccount": "5020 - Direct Materials",
                "glSuggestions": [],
            }
        return {
            "detectedType": "Sales Summary",
            "confidenceScore": 0.95,
            "target": "WP1",
            "evidence": ["Detected sales-summary style columns in uploaded file."],
            "warnings": [],
            "suggestedGlAccount": "4100 - Sales Revenue",
            "glSuggestions": [],
        }

    revenue_detection = None
    if lane in {"auto", "sales"}:
        revenue_detection = classify_revenue_context(
            filename=filename,
            headers=headers,
            sample_rows=sample_rows,
            row_count=row_count,
        )
    if revenue_detection:
        return revenue_detection

    return {}


def extract_document_party(filename: str, text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    ignore_tokens = {
        "tax invoice", "invoice", "invoice no", "invoice number", "invoice date", "statement", "bank statement",
        "official receipt", "receipt", "date", "reference", "page 1", "page 2",
    }
    preferred_lines: list[str] = []
    for line in lines[:16]:
        lowered = normalize(line)
        if not lowered or len(lowered) < 4:
            continue
        if any(token in lowered for token in ignore_tokens):
            continue
        if any(token in lowered for token in ("phone", "tel", "fax", "email", "company no", "tin no", "sst no", "reg no")):
            continue
        if AMOUNT_TOKEN_PATTERN.search(line):
            continue
        if sum(1 for character in line if character.isalpha()) < 3:
            continue
        preferred_lines.append(line)

    for line in preferred_lines:
        lowered = normalize(line)
        if any(token in lowered for token in COMPANY_SUFFIXES):
            preferred_lines.insert(0, line)
            break

    ranked_candidates: list[tuple[int, str]] = []
    seen_cleaned: set[str] = set()
    raw_candidates = [
        _extract_invoice_header_company(text),
        *_extract_company_candidates(text[:2500]),
        _extract_labeled_value(LABELED_PARTY_PATTERN, text[:2500]),
        *preferred_lines[:4],
    ]
    for raw_candidate in raw_candidates:
        candidate = normalize_whitespace(raw_candidate)
        cleaned = clean_party_value(candidate)
        if not cleaned:
            continue
        normalized_cleaned = normalize(cleaned)
        if normalized_cleaned in seen_cleaned:
            continue
        seen_cleaned.add(normalized_cleaned)

        ranked_candidates.append((_score_party_candidate(cleaned), cleaned))

    if ranked_candidates:
        ranked_candidates.sort(key=lambda item: item[0], reverse=True)
        return ranked_candidates[0][1][:80]

    return _party_from_filename(filename)


def looks_like_bank_statement_text(filename: str, text: str) -> bool:
    content = normalize(f"{filename} {text[:5000]}")
    if any(token in content for token in ("bank statement", "account statement", "opening balance", "closing balance", "beginning balance", "ending balance")):
        return True
    if any(token in content for token in ("cimb", "maybank", "public bank", "rhb", "hong leong", "uob", "ocbc", "ambank", "bsn")) and "statement" in content:
        return True
    transaction_lines = 0
    for line in text.splitlines()[:200]:
        normalized_line = normalize(line)
        if not normalized_line:
            continue
        if BANK_TEXT_LINE_PATTERN.match(line) and len(AMOUNT_TOKEN_PATTERN.findall(line)) >= 2:
            transaction_lines += 1
        if transaction_lines >= 3:
            return True
    return False


def extract_bank_statement_rows_from_text(
    *,
    filename: str,
    content_type: str,
    data: bytes,
    text: str,
    client_entity_name: str | None = None,
) -> list[dict[str, Any]]:
    if not looks_like_bank_statement_text(filename, text):
        return []

    uploaded_at = current_timestamp()
    items: list[dict[str, Any]] = []
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    def normalize_bank_date(raw_value: str) -> str:
        normalized = normalize_date(raw_value)
        if normalized:
            return normalized
        short = re.match(r"^\s*(\d{1,2})[/-](\d{1,2})\s*$", raw_value)
        if short:
            return f"{int(short.group(1)):02d} {month_name(int(short.group(2)))}"
        return ""

    for index, line in enumerate(lines):
        match = BANK_TEXT_LINE_PATTERN.match(line)
        if not match:
            continue

        date = normalize_bank_date(match.group("date"))
        if not date:
            continue

        amount_matches = [
            found for found in AMOUNT_TOKEN_PATTERN.finditer(match.group("rest"))
            if "." in found.group(0)
        ]
        if len(amount_matches) < 2:
            continue

        rest = match.group("rest")
        description = rest[: amount_matches[0].start()].strip(" -:")
        leading_amount_layout = amount_matches[0].start() <= 2
        if leading_amount_layout and len(amount_matches) >= 2:
            description = rest[amount_matches[1].end() :].strip(" -:")
            if not description:
                continuation_parts: list[str] = []
                look_ahead = index + 1
                while look_ahead < len(lines) and len(continuation_parts) < 3:
                    continuation = lines[look_ahead].strip(" -:")
                    if not continuation:
                        look_ahead += 1
                        continue
                    if BANK_TEXT_LINE_PATTERN.match(continuation):
                        break
                    normalized_continuation = normalize(continuation)
                    if normalized_continuation.startswith("balance c/f") or normalized_continuation.startswith("balance b/f"):
                        break
                    if len(AMOUNT_TOKEN_PATTERN.findall(continuation)) >= 2 and not any(character.isalpha() for character in continuation):
                        break
                    continuation_parts.append(continuation)
                    look_ahead += 1
                description = " ".join(continuation_parts).strip()
        if not description or len(description) < 2:
            continue

        amounts = [parse_amount(found.group(0)) for found in amount_matches]
        non_zero_amounts = [value for value in amounts if value > 0]
        if len(non_zero_amounts) < 2:
            continue

        transaction_amount = non_zero_amounts[0] if leading_amount_layout else (non_zero_amounts[-2] if len(non_zero_amounts) >= 2 else non_zero_amounts[0])
        lowered = normalize(line)
        money_in = 0.0
        money_out = 0.0
        if any(token in lowered for token in (" credit ", " cr ", "deposit", "money in", "received", "transfer in", "dep-merchant", "dep merchant", "duitnow qr cr")):
            money_in = transaction_amount
        elif any(token in lowered for token in (" debit ", " dr ", "withdrawal", "payment", "charges", "bank charge", "cash out")):
            money_out = transaction_amount
        else:
            money_out = transaction_amount

        reference = extract_reference(line, filename)
        items.append(
            build_item(
                file_name=f"{filename} row {len(items) + 1}",
                file_type=content_type,
                file_size=len(data),
                uploaded_at=uploaded_at,
                client_entity_name=client_entity_name,
                detected_type="Bank Statement",
                target="WP2",
                date=date,
                reference=reference,
                party=description[:120],
                amount=transaction_amount,
                money_in=money_in,
                money_out=money_out,
                suggested_gl_account="",
                notes="Extracted from PDF/text bank statement rows.",
                evidence=["Detected bank statement line pattern from readable text."],
                warnings=[] if money_in or money_out else ["Bank transaction direction was inferred conservatively. Review in WP2."],
                raw_preview=line[:240],
                overall_confidence_score=0.82 if money_in or money_out else 0.72,
                extraction_method="bank-text-parser",
            )
        )
        if len(items) >= 200:
            break

    if items:
        return items

    vertical_date_pattern = re.compile(r"^\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?$")
    amount_only_pattern = re.compile(r"^\(?(?:(?:\d{1,3}(?:,\d{3})+)|\d+)\.\d{2}\)?$")

    def is_balance_marker(value: str) -> bool:
        lowered = normalize(value)
        return (
            lowered.startswith("balance from last statement")
            or lowered.startswith("balance b f")
            or lowered.startswith("balance c f")
            or lowered.startswith("closing balance in this statement")
        )

    def is_layout_noise(value: str) -> bool:
        lowered = normalize(value)
        if not lowered:
            return True
        if lowered in {"tarikh", "urus niaga", "debit", "kredit", "baki", "date", "transaction", "credit", "balance"}:
            return True
        return any(
            token in lowered
            for token in (
                "page ",
                "muka surat",
                "account number",
                "statement date",
                "account type",
                "penyata ini dicetak",
                "this is a computer generated statement",
                "protected by pidm",
                "dilindungi oleh pidm",
            )
        )

    def classify_bank_direction(description: str) -> tuple[float, float]:
        lowered = f" {normalize(description)} "
        if any(token in lowered for token in (" credit ", " cr ", " deposit ", " money in ", " received ", " transfer in ", " dep merchant ", " dep merchant pymt ", " duitnow qr cr ", " rpp cr ")):
            return 1.0, 0.0
        if any(token in lowered for token in (" debit ", " dr ", " withdrawal ", " payment ", " charges ", " bank charge ", " cash out ", " trsf dr ")):
            return 0.0, 1.0
        return 0.0, 1.0

    current_date = ""
    pending_description: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        normalized_line = normalize(line)

        if vertical_date_pattern.match(line):
            current_date = normalize_bank_date(line)
            pending_description = []
            index += 1
            continue

        if not current_date:
            index += 1
            continue

        if is_balance_marker(line):
            if normalized_line.startswith("balance c/f"):
                current_date = ""
                pending_description = []
            index += 1
            if index < len(lines) and amount_only_pattern.match(lines[index]):
                index += 1
            continue

        if is_layout_noise(line):
            index += 1
            continue

        if (
            amount_only_pattern.match(line)
            and index + 1 < len(lines)
            and amount_only_pattern.match(lines[index + 1])
        ):
            transaction_amount = parse_amount(line)
            index += 2
            description_parts = [*pending_description]
            pending_description = []

            while index < len(lines):
                candidate = lines[index].strip(" -:")
                if not candidate:
                    index += 1
                    continue
                if vertical_date_pattern.match(candidate) or is_balance_marker(candidate):
                    break
                if (
                    amount_only_pattern.match(candidate)
                    and index + 1 < len(lines)
                    and amount_only_pattern.match(lines[index + 1])
                ):
                    break
                if is_layout_noise(candidate):
                    index += 1
                    continue
                description_parts.append(candidate)
                index += 1

            description = normalize_whitespace(" ".join(description_parts))
            if not description or transaction_amount <= 0:
                continue

            money_in_factor, money_out_factor = classify_bank_direction(description)
            money_in = transaction_amount * money_in_factor
            money_out = transaction_amount * money_out_factor
            reference = extract_reference(description, filename)
            items.append(
                build_item(
                    file_name=f"{filename} row {len(items) + 1}",
                    file_type=content_type,
                    file_size=len(data),
                    uploaded_at=uploaded_at,
                    client_entity_name=client_entity_name,
                    detected_type="Bank Statement",
                    target="WP2",
                    date=current_date,
                    reference=reference,
                    party=description[:120],
                    amount=transaction_amount,
                    money_in=money_in,
                    money_out=money_out,
                    suggested_gl_account="",
                    notes="Extracted from vertically laid out PDF/text bank statement rows.",
                    evidence=["Detected vertically split bank statement row pattern from readable text."],
                    warnings=[] if money_in or money_out else ["Bank transaction direction was inferred conservatively. Review in WP2."],
                    raw_preview=f"{current_date} {description}"[:240],
                    overall_confidence_score=0.84 if money_in or money_out else 0.72,
                    extraction_method="bank-text-parser",
                )
            )
            if len(items) >= 600:
                break
            continue

        pending_description.append(line)
        if len(pending_description) > 4:
            pending_description = pending_description[-4:]
        index += 1

    return items


def _file_channel_label(filename: str) -> str:
    match = re.search(r"\(([^)]+)\)", filename)
    if match:
        return match.group(1).strip()

    cleaned = re.sub(r"\.[^.]+$", "", filename)
    for token in re.split(r"[_\-\s]+", cleaned):
        candidate = token.strip()
        if candidate and candidate.lower() not in {"sales", "by", "category", "daily", "report"}:
            if any(character.isalpha() for character in candidate):
                return candidate
    return ""


def _looks_like_weak_sales_label(value: str) -> bool:
    normalized = normalize(value)
    if not normalized:
        return True
    if normalized in GENERIC_SALES_LABELS:
        return True
    if normalized in DAY_OR_TIME_TOKENS:
        return True
    if re.fullmatch(r"(20\d{2}|\d+(?:\.\d+)?)", normalized):
        return True

    tokens = normalized.split()
    if tokens and all(token in DAY_OR_TIME_TOKENS or re.fullmatch(r"\d+(?:\.\d+)?", token) for token in tokens):
        return True

    return False


def _sales_summary_party(
    *,
    filename: str,
    headers: list[str],
    cells: list[str],
) -> tuple[str, float, list[str], list[str]]:
    preferred_columns = [
        "category",
        "item",
        "product",
        "menu item",
        "menu",
        "outlet",
        "counter",
        "department",
        "class",
        "segment",
        "description",
    ]

    selected_value = ""
    selected_header = ""
    for candidate in preferred_columns:
        for index, header in enumerate(headers):
            if candidate in header and index < len(cells):
                value = cells[index].strip()
                if value:
                    selected_value = value
                    selected_header = header
                    break
        if selected_value:
            break

    channel = _file_channel_label(filename)
    warnings: list[str] = []
    evidence: list[str] = []

    if selected_value:
        evidence.append(f"Sales label came from `{selected_header}` column.")
        if _looks_like_weak_sales_label(selected_value):
            if channel and normalize(channel) != normalize(selected_value):
                return f"{channel} - {selected_value}", 0.68, [*evidence, "Prepended file channel to strengthen a weak row label."], [
                    "Sales row label looked weak, so the file channel was prepended. Review before posting."
                ]
            return selected_value, 0.62, evidence, ["Sales row label looked weak. Review before posting."]

        if channel and normalize(channel) not in normalize(selected_value):
            return f"{channel} - {selected_value}", 0.92, [*evidence, "Prepended file channel for clearer BK-facing description."], warnings
        return selected_value, 0.9, evidence, warnings

    if channel:
        return channel, 0.58, ["Used the file channel because no strong category/item column was found."], [
            "Sales row description fell back to the file channel. Review before posting."
        ]

    return "", 0.0, [], ["Sales row description could not be determined."]


def item_from_cells(
    *,
    filename: str,
    content_type: str,
    file_size: int,
    uploaded_at: str,
    headers: list[str],
    raw_headers: list[str],
    cells: list[str],
    row_text: str,
    extraction_method: str,
    context_detection: dict[str, Any] | None = None,
    upload_lane: str | None = None,
    client_entity_name: str | None = None,
    source_file_name: str | None = None,
    header_row_number: int | None = None,
    source_row_number: int | None = None,
    raw_context_rows: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    header_text = " ".join(headers)
    detection = context_detection or detect_type(f"{filename} {header_text} {row_text}", upload_lane)

    date = cell_at(headers, cells, ["date", "transaction date"]) or normalize_date(row_text)
    if detection["detectedType"] == "Sales Summary":
        description, party_confidence, party_evidence, party_warnings = _sales_summary_party(
            filename=filename,
            headers=headers,
            cells=cells,
        )
    else:
        description = cell_prefer(headers, cells, ["description", "bank description", "vendor", "customer", "party", "category", "outlet", "item", "product"])
        party_confidence = 0.9 if description else 0.0
        party_evidence = []
        party_warnings = []
    reference = cell_at(headers, cells, ["reference", "doc ref", "document ref", "invoice", "ref"])
    type_from_row = cell_at(headers, cells, ["document type", "doc type", "type"])
    amount_text = cell_prefer(headers, cells, ["net sales", "total sales", "amount", "total", "gross sales", "sales"])
    money_in = parse_amount(cell_at(headers, cells, ["money in", "deposit", "credit"]))
    money_out = parse_amount(cell_at(headers, cells, ["money out", "withdrawal", "debit"]))
    signed_amount, signed_direction = parse_signed_amount(amount_text)
    amount = max(signed_amount, money_in, money_out)

    if detection["target"] == "WP2" and not money_in and not money_out and amount:
        if signed_direction == "OUT":
            money_out = amount
        else:
            money_in = amount

    detected_type = type_from_row if type_from_row in {
        "Sales Invoice",
        "Sales Summary",
        "Purchase Invoice",
        "Payment Voucher",
        "Receipt",
        "Payroll Summary",
        "Loan / HP Statement",
        "Merchant Statement",
        "Utility Bill",
        "Bank Statement",
    } else detection["detectedType"]

    invoice_role, role_evidence, role_warnings = infer_invoice_role_from_text(
        filename=filename,
        text=f"{header_text}\n{row_text}",
        client_entity_name=client_entity_name or "",
    )
    if invoice_role == "purchase":
        detected_type = "Purchase Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "5020 - Direct Materials"
    elif invoice_role == "sales":
        detected_type = "Sales Invoice"
        detection["target"] = "WP1"
        detection["suggestedGlAccount"] = "4100 - Sales Revenue"

    completeness_score = 0.0
    completeness_score += 0.2 if date else 0
    completeness_score += 0.2 if amount > 0 else 0
    completeness_score += 0.15 if description else 0
    completeness_score += 0.15 if reference else 0
    completeness_score += 0.15 if detected_type != "Unknown" else 0
    completeness_score += 0.15 if cell_at(headers, cells, ["gl account", "account"]) or detection["suggestedGlAccount"] else 0

    warnings = [
        *detection["warnings"],
        *role_warnings,
        *party_warnings,
        "" if date else "Date was not detected.",
        "" if amount > 0 else "Amount was not detected.",
    ]

    if detected_type == "Sales Summary" and amount <= 0:
        return None

    return build_item(
        file_name=filename,
        file_type=content_type,
        file_size=file_size,
        uploaded_at=uploaded_at,
        client_entity_name=client_entity_name,
        detected_type=detected_type,
        target=detection["target"],
        date=date,
        reference=reference,
        party=description,
        amount=amount,
        money_in=money_in,
        money_out=money_out,
        suggested_gl_account=cell_at(headers, cells, ["gl account", "account"]) or detection["suggestedGlAccount"],
        notes=cell_at(headers, cells, ["note", "notes", "remarks"]) or "Extracted from uploaded workbook/spreadsheet.",
        evidence=[*detection["evidence"], *role_evidence],
        warnings=[warning for warning in warnings if warning],
        raw_preview=row_text[:240],
        overall_confidence_score=max(detection["confidenceScore"], completeness_score),
        extraction_method=extraction_method,
        gl_suggestions=detection.get("glSuggestions"),
        field_confidence={
            "date": 0.95 if date else 0.0,
            "reference": 0.92 if reference else 0.0,
            "party": party_confidence,
            "amount": 0.96 if amount > 0 else 0.0,
            "suggestedGlAccount": 0.88 if cell_at(headers, cells, ["gl account", "account"]) or detection["suggestedGlAccount"] else 0.0,
        },
        field_evidence={
            "party": party_evidence,
        },
        raw_source={
            "file_name": source_file_name or filename,
            "header_row_number": header_row_number,
            "row_number": source_row_number,
            "headers": headers,
            "display_headers": raw_headers,
            "cells": cells,
            "context_rows": raw_context_rows or [],
        },
    )

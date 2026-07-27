import openpyxl, sys
from pathlib import Path

path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("wp_audit_small.xlsx")
wb = openpyxl.load_workbook(path)
ws = wb.active
headers = [c.value for c in ws[1]]
for row in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(headers, row))
    print(f"File:    {d['Source File']}")
    print(f"Vendor:  {d['Vendor / Customer']}")
    print(f"Date:    {d['Date']}")
    print(f"Amount:  {d['Amount']}")
    print(f"Type:    {d['Document Type']}")
    print(f"Status:  {d['Status']}")
    print()

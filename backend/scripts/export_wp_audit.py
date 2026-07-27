"""
Extract all documents in a folder and export results to WP1 / WP2 Excel sheets.

Usage (from repo root):
    python backend/scripts/export_wp_audit.py
    python backend/scripts/export_wp_audit.py --client "Your Entity Sdn Bhd"
    python backend/scripts/export_wp_audit.py --max-files 20   # quick sample run
    python backend/scripts/export_wp_audit.py --lane purchases  # bias toward purchase invoices
    python backend/scripts/export_wp_audit.py --provider openai --api-key sk-...
    python backend/scripts/export_wp_audit.py --provider openrouter --model deepseek/deepseek-vl2
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

# ── Path setup ────────────────────────────────────────────────────────────────
BACKEND_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = BACKEND_ROOT.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))


# ── Load .env.production.local before importing backend modules ───────────────
def _load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        match = re.match(r"^([A-Z0-9_]+)=(.*)$", line)
        if not match:
            continue
        key, raw = match.group(1), match.group(2)
        if len(raw) >= 2 and raw[0] == raw[-1] and raw[0] in ('"', "'"):
            raw = raw[1:-1]
        # Vercel CLI appends literal \r\n inside quoted values — strip them
        value = raw.replace("\\r\\n", "").replace("\\r", "").replace("\\n", "").strip()
        if value:
            os.environ[key] = value


_load_env_file(REPO_ROOT / ".env.production.local")

# ── Backend imports ───────────────────────────────────────────────────────────
from app.extractors import extract_intake_items  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# ── Styling constants ─────────────────────────────────────────────────────────
NAVY_FILL   = PatternFill("solid", fgColor="1F3864")
GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT   = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
)

CONTENT_TYPES: dict[str, str] = {
    ".pdf":  "application/pdf",
    ".csv":  "text/csv",
    ".txt":  "text/plain",
    ".tsv":  "text/tab-separated-values",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".xls":  "application/vnd.ms-excel",
    ".png":  "image/png",
    ".jpg":  "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}

# WP1 columns: match the template exactly, then add audit helpers at the end
WP1_HEADERS = [
    "Date",
    "Document Ref",
    "Vendor / Customer",
    "Document Type",
    "Amount",
    "GL Account",
    "Notes",
    "Status",          # Accepted / Needs Review
    "Source File",     # original filename for traceability
]
WP1_WIDTHS = [12, 22, 32, 22, 12, 28, 42, 15, 38]

# WP2 columns: match the template exactly, then audit helpers
WP2_HEADERS = [
    "Date",
    "Bank Description",
    "Reference",
    "Money In",
    "Money Out",
    "Notes",
    "Status",
    "Source File",
]
WP2_WIDTHS = [12, 32, 22, 12, 12, 42, 15, 38]


# ── Helpers ───────────────────────────────────────────────────────────────────

def _header_row(ws: openpyxl.worksheet.worksheet.Worksheet, headers: list[str]) -> None:
    for col, heading in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=heading)
        cell.fill = NAVY_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28


def _body_cell(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row: int,
    col: int,
    value: object,
    fill: PatternFill,
    wrap: bool = False,
) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = BODY_FONT
    cell.fill = fill
    cell.border = THIN_BORDER
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)


def _notes_for(item: dict) -> str:
    parts: list[str] = []
    if item.get("notes"):
        parts.append(str(item["notes"]).strip())
    review_signals = (
        "not detected",
        "inferred",
        "needs review",
        "weak",
        "could not",
        "mismatch",
    )
    for w in item.get("warnings", []):
        low = w.lower()
        if any(s in low for s in review_signals):
            parts.append(w)
            if len(parts) >= 3:
                break
    return "; ".join(parts)[:240]


# ── Core builder ─────────────────────────────────────────────────────────────

def build_workbook(
    folder: Path,
    *,
    client_entity_name: str,
    upload_lane: str,
    max_files: int,
    verbose: bool,
) -> tuple[openpyxl.Workbook, dict]:
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "WP1 Document Ledger"
    _header_row(ws1, WP1_HEADERS)
    ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet("WP2 Bank Verification")
    _header_row(ws2, WP2_HEADERS)
    ws2.freeze_panes = "A2"

    wp1_row = 2
    wp2_row = 2

    stats = {
        "files_processed": 0,
        "files_errored": 0,
        "wp1_items": 0,
        "wp2_items": 0,
        "accepted": 0,
        "needs_review": 0,
    }

    all_files = sorted(
        p for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in CONTENT_TYPES
    )
    if max_files > 0:
        all_files = all_files[:max_files]

    total = len(all_files)

    for idx, path in enumerate(all_files, 1):
        if verbose:
            label = f"[{idx:>3}/{total}] {path.name}"
            print(f"  {label:<60}", end="", flush=True)

        content_type = CONTENT_TYPES[path.suffix.lower()]
        try:
            items = extract_intake_items(
                path.name,
                content_type,
                path.read_bytes(),
                upload_lane=upload_lane,
                client_entity_name=client_entity_name,
            )
        except Exception as exc:
            if verbose:
                print(f"  ERROR: {exc}")
            stats["files_errored"] += 1
            stats["files_processed"] += 1
            continue

        stats["files_processed"] += 1
        n_accepted = sum(1 for i in items if i.get("status") == "Accepted")
        n_review   = sum(1 for i in items if i.get("status") != "Accepted")
        stats["accepted"]     += n_accepted
        stats["needs_review"] += n_review

        if verbose:
            print(f"  {len(items)} item(s) - {n_accepted} ok  {n_review} review")

        for item in items:
            target = item.get("target", "WP1")
            status = item.get("status", "Needs Review")
            fill   = GREEN_FILL if status == "Accepted" else YELLOW_FILL
            notes  = _notes_for(item)

            if target == "WP2":
                stats["wp2_items"] += 1
                row_vals = [
                    item.get("date") or "",
                    item.get("party") or "",
                    item.get("reference") or "",
                    item.get("moneyIn")  or "",
                    item.get("moneyOut") or "",
                    notes,
                    status,
                    path.name,
                ]
                for col, val in enumerate(row_vals, 1):
                    _body_cell(ws2, wp2_row, col, val, fill, wrap=(col in (2, 6)))
                wp2_row += 1

            else:
                stats["wp1_items"] += 1
                row_vals = [
                    item.get("date") or "",
                    item.get("reference") or "",
                    item.get("party") or "",
                    item.get("detectedType") or "",
                    item.get("amount")  or "",
                    item.get("suggestedGlAccount") or "",
                    notes,
                    status,
                    path.name,
                ]
                for col, val in enumerate(row_vals, 1):
                    _body_cell(ws1, wp1_row, col, val, fill, wrap=(col in (3, 7)))
                wp1_row += 1

    # Column widths
    for col, width in enumerate(WP1_WIDTHS, 1):
        ws1.column_dimensions[ws1.cell(row=1, column=col).column_letter].width = width

    for col, width in enumerate(WP2_WIDTHS, 1):
        ws2.column_dimensions[ws2.cell(row=1, column=col).column_letter].width = width

    # Auto-filter on both sheets
    ws1.auto_filter.ref = ws1.dimensions
    ws2.auto_filter.ref = ws2.dimensions

    return wb, stats


# ── Entry point ───────────────────────────────────────────────────────────────

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract all test docs and export to WP1 / WP2 Excel.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "--folder",
        default=str(REPO_ROOT / "test docs"),
        help="Folder containing source documents (default: test docs/).",
    )
    parser.add_argument(
        "--client",
        default="",
        help="Client entity name — used to determine purchase vs. sales role.",
    )
    parser.add_argument(
        "--lane",
        default="auto",
        choices=["auto", "purchases", "sales", "bank", "payments"],
        help="Upload lane hint (default: auto).",
    )
    parser.add_argument(
        "--max-files",
        type=int,
        default=0,
        help="Limit to first N files — useful for a quick sample run (0 = all).",
    )
    parser.add_argument(
        "--output",
        default=str(REPO_ROOT / "wp_audit.xlsx"),
        help="Output Excel file path (default: wp_audit.xlsx in repo root).",
    )
    parser.add_argument(
        "--provider",
        default="",
        choices=["", "anthropic", "openai", "openrouter", "gemini"],
        help="Override AI provider (anthropic/openai/openrouter). Defaults to .env setting.",
    )
    parser.add_argument(
        "--model",
        default="",
        help="Override vision model (e.g. gpt-4o, claude-sonnet-4-6, deepseek/deepseek-vl2).",
    )
    parser.add_argument(
        "--api-key",
        default="",
        help="API key for the chosen provider. Sets OPENAI_API_KEY / ANTHROPIC_API_KEY as needed.",
    )
    args = parser.parse_args()

    # Apply provider/model/key overrides before any extraction calls
    if args.provider:
        os.environ["AI_PROVIDER"] = args.provider
    provider = os.environ.get("AI_PROVIDER", "").strip().lower() or "openai"
    if args.api_key:
        key_env = {
            "anthropic":  "ANTHROPIC_API_KEY",
            "openai":     "OPENAI_API_KEY",
            "openrouter": "OPENROUTER_API_KEY",
            "gemini":     "GEMINI_API_KEY",
        }.get(provider, "OPENAI_API_KEY")
        os.environ[key_env] = args.api_key
    if args.model:
        model_env = {
            "anthropic":  "ANTHROPIC_VISION_MODEL",
            "openai":     "OPENAI_VISION_MODEL",
            "openrouter": "OPENROUTER_VISION_MODEL",
            "gemini":     "GEMINI_VISION_MODEL",
        }.get(provider, "OPENAI_VISION_MODEL")
        os.environ[model_env] = args.model

    # Lane-based vision routing (only when not explicitly overridden via --provider/--model)
    if not args.provider and not args.model:
        if args.lane in ("bank", "sales"):
            os.environ["OPENAI_VISION_ENABLED"] = "false"
        else:
            os.environ["OPENAI_VISION_ENABLED"] = "true"

    folder = Path(args.folder)
    if not folder.is_dir():
        print(f"ERROR: folder not found: {folder}")
        return 1

    all_files = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in CONTENT_TYPES]
    target_count = min(len(all_files), args.max_files) if args.max_files > 0 else len(all_files)

    from app.extractors.ai_provider import configured_provider, vision_model
    from app.extractors.vision import can_use_vision
    print("\n" + "=" * 64)
    print(f"  WP Audit -- extracting {target_count} of {len(all_files)} file(s)")
    print(f"  Folder   : {folder}")
    print(f"  Client   : {args.client or '(not set)'}")
    print(f"  Lane     : {args.lane}")
    if can_use_vision():
        print(f"  Provider : {configured_provider()}")
        print(f"  Model    : {vision_model()}")
    else:
        print(f"  Vision   : off (text extraction only)")
    print(f"  Output   : {args.output}")
    print("=" * 64 + "\n")

    wb, stats = build_workbook(
        folder,
        client_entity_name=args.client,
        upload_lane=args.lane,
        max_files=args.max_files,
        verbose=True,
    )

    output_path = Path(args.output)
    wb.save(output_path)

    print(f"\n{'='*64}")
    print(f"  Files processed : {stats['files_processed']}")
    if stats["files_errored"]:
        print(f"  Files errored   : {stats['files_errored']}")
    print(f"  WP1 items       : {stats['wp1_items']}")
    print(f"  WP2 items       : {stats['wp2_items']}")
    print(f"  Accepted        : {stats['accepted']}")
    print(f"  Needs Review    : {stats['needs_review']}")
    print(f"\n  Saved to: {output_path}")
    print("=" * 64 + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
